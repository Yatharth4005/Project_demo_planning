from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import BTRUpload, CastingPlan, RuleMaster


def sample_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Section", "Steel Grade", "Order Qty", "Order Number", "Due Date"])
    sheet.append(["Rail 60E1", "R260", 128, "SO-001", "2025-06-03"])
    sheet.append(["Rail 60E1", "R260", 72, "SO-002", "2025-06-04"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class PlanningWorkflowTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="planner", password="safe-test-password"
        )
        RuleMaster.objects.create(
            section="Rail 60E1",
            grade="R260",
            lfa_batch_tonnes=128,
            target_yield_percent=93.2,
        )
        self.client.force_login(self.user)

    def test_complete_btr_to_export_workflow(self):
        response = self.client.post(
            reverse("btr_upload"),
            {
                "workbook": SimpleUploadedFile(
                    "btr.xlsx",
                    sample_workbook(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            },
        )
        upload = BTRUpload.objects.get()
        self.assertRedirects(
            response, reverse("campaign_create", args=[upload.pk])
        )
        self.assertEqual(upload.row_count, 2)
        self.assertEqual(float(upload.total_demand_tonnes), 200)

        response = self.client.post(
            reverse("campaign_create", args=[upload.pk]),
            {"material": "Rail 60E1|||R260", "strategy": "yield"},
        )
        plan = CastingPlan.objects.get()
        self.assertRedirects(response, reverse("casting_detail", args=[plan.pk]))
        self.assertGreater(len(plan.heat_sequence), 1)

        response = self.client.get(reverse("export_plan", args=[plan.pk]))
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        exported = load_workbook(BytesIO(content))
        self.assertEqual(exported["Casting Plan"]["B2"].value, plan.plan_number)

    def test_workflow_requires_authentication(self):
        self.client.logout()
        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, f"{reverse('login')}?next=/")
