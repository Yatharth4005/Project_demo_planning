from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

import pandas as pd
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import BTRUpload, CampaignPlan, CastingPlan, RuleMaster


COLUMN_ALIASES = {
    "section": {"section", "profile", "material", "section name"},
    "grade": {"grade", "steel grade", "quality"},
    "quantity": {"quantity", "qty", "demand", "order qty", "tonnage", "weight"},
    "order": {"order", "order no", "order number", "sales order"},
    "due_date": {"due date", "delivery date", "required date"},
}


def _normalize_column(value):
    return " ".join(str(value).strip().lower().replace("_", " ").split())


def parse_btr(upload: BTRUpload) -> BTRUpload:
    extension = Path(upload.workbook.name).suffix.lower()
    engine = "openpyxl" if extension in {".xlsx", ".xlsm"} else None
    try:
        frame = pd.read_excel(upload.workbook.path, engine=engine)
        columns = {_normalize_column(column): column for column in frame.columns}
        resolved = {}
        for field, aliases in COLUMN_ALIASES.items():
            match = next((columns[name] for name in aliases if name in columns), None)
            if match is not None:
                resolved[field] = match

        missing = [field for field in ("section", "grade", "quantity") if field not in resolved]
        if missing:
            raise ValueError(
                "Missing required columns: " + ", ".join(field.title() for field in missing)
            )

        normalized = pd.DataFrame()
        normalized["section"] = frame[resolved["section"]].astype(str).str.strip()
        normalized["grade"] = frame[resolved["grade"]].astype(str).str.strip()
        normalized["quantity"] = pd.to_numeric(
            frame[resolved["quantity"]], errors="coerce"
        )
        normalized["order"] = (
            frame[resolved["order"]].astype(str)
            if "order" in resolved
            else [f"ROW-{index + 2}" for index in range(len(frame))]
        )
        normalized["due_date"] = (
            pd.to_datetime(frame[resolved["due_date"]], errors="coerce")
            if "due_date" in resolved
            else pd.NaT
        )
        normalized = normalized.dropna(subset=["quantity"])
        normalized = normalized[
            (normalized["section"] != "") & (normalized["grade"] != "")
        ]
        normalized = normalized[normalized["quantity"] > 0]
        if normalized.empty:
            raise ValueError("No valid positive-demand rows were found.")

        records = []
        for row in normalized.to_dict("records"):
            due_date = row["due_date"]
            records.append(
                {
                    "section": row["section"],
                    "grade": row["grade"],
                    "quantity": float(row["quantity"]),
                    "order": row["order"],
                    "due_date": (
                        due_date.date().isoformat() if not pd.isna(due_date) else None
                    ),
                }
            )
        upload.status = BTRUpload.Status.VALIDATED
        upload.row_count = len(records)
        upload.total_demand_tonnes = Decimal(
            str(sum(row["quantity"] for row in records))
        ).quantize(Decimal("0.01"))
        upload.normalized_orders = records
        upload.validation_errors = []
    except Exception as error:
        upload.status = BTRUpload.Status.FAILED
        upload.validation_errors = [str(error)]
        upload.normalized_orders = []
    upload.save()
    return upload


def materials_for(upload: BTRUpload):
    return sorted(
        {
            (row["section"], row["grade"])
            for row in upload.normalized_orders
        }
    )


def material_summary(upload: BTRUpload, section: str, grade: str):
    orders = [
        row
        for row in upload.normalized_orders
        if row["section"] == section and row["grade"] == grade
    ]
    return {
        "demand": round(sum(row["quantity"] for row in orders), 2),
        "orders": len(orders),
        "dates": sorted(row["due_date"] for row in orders if row["due_date"]),
    }


@transaction.atomic
def create_casting_plan(upload, user, section, grade, strategy):
    rows = [
        row
        for row in upload.normalized_orders
        if row["section"] == section and row["grade"] == grade
    ]
    if strategy == CampaignPlan.Strategy.DELIVERY:
        rows.sort(key=lambda row: (row["due_date"] or "9999-12-31", row["order"]))
    else:
        rows.sort(key=lambda row: (-row["quantity"], row["order"]))

    demand = Decimal(str(sum(row["quantity"] for row in rows))).quantize(
        Decimal("0.01")
    )
    rule, _ = RuleMaster.objects.get_or_create(section=section, grade=grade)
    target_yield = Decimal(rule.target_yield_percent)
    total_cast = (demand / (target_yield / Decimal("100"))).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )
    campaign = CampaignPlan.objects.create(
        btr=upload,
        created_by=user,
        section=section,
        grade=grade,
        strategy=strategy,
        demand_tonnes=demand,
        order_count=len(rows),
        sequence=rows,
    )

    batch_size = Decimal(rule.lfa_batch_tonnes)
    remaining = total_cast
    heats = []
    index = 1
    while remaining > 0:
        cast_weight = min(batch_size, remaining)
        order_load = (cast_weight * target_yield / Decimal("100")).quantize(
            Decimal("0.01")
        )
        heats.append(
            {
                "sequence": index,
                "batch": f"LFA-{timezone.now():%y%m}{index:03d}",
                "grade": grade,
                "billets": max(1, int(cast_weight / Decimal("1.32"))),
                "cast_weight": float(cast_weight),
                "order_load": float(order_load),
                "yield": float(target_yield),
            }
        )
        remaining -= cast_weight
        index += 1

    count = CastingPlan.objects.count() + 1
    return CastingPlan.objects.create(
        campaign=campaign,
        plan_number=f"CP-{timezone.now():%Y}-{count:04d}",
        total_cast_tonnes=total_cast,
        expected_yield_percent=target_yield,
        heat_sequence=heats,
    )


def campaign_workbook(plan: CastingPlan) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Casting Plan"
    sheet.append(["CASTING PLAN OPTIMIZER"])
    sheet.append(["Plan ID", plan.plan_number])
    sheet.append(["Section", plan.campaign.section])
    sheet.append(["Grade", plan.campaign.grade])
    sheet.append(["Total cast (T)", float(plan.total_cast_tonnes)])
    sheet.append(["Expected yield (%)", float(plan.expected_yield_percent)])
    sheet.append([])
    headers = [
        "Sequence",
        "Heat / Batch",
        "Grade",
        "Billets",
        "Cast Weight (T)",
        "Order Load (T)",
        "Yield (%)",
    ]
    sheet.append(headers)
    for heat in plan.heat_sequence:
        sheet.append(
            [
                heat["sequence"],
                heat["batch"],
                heat["grade"],
                heat["billets"],
                heat["cast_weight"],
                heat["order_load"],
                heat["yield"],
            ]
        )
    for cell in sheet[8]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D34")
    for width, column in zip([12, 18, 16, 12, 18, 18, 14], "ABCDEFG"):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A9"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
